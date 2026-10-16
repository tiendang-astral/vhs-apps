"""Kiểm thử toàn bộ quy trình qua HTTP (chạy server thật trên SQLite tạm).
Chạy:  python tests/test_flow.py"""
import os
import socket
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import requests

ROOT = Path(__file__).resolve().parent.parent
PASS, FAIL = [], []


def check(name, cond, info=''):
    (PASS if cond else FAIL).append(name)
    print(('  ✓ ' if cond else '  ✗ ') + name + ('' if cond else f'  → {info}'))


class Client:
    def __init__(self, base):
        self.base, self.tok = base, None

    def req(self, method, path, **kw):
        h = {'Authorization': f'Bearer {self.tok}'} if self.tok else {}
        r = requests.request(method, self.base + path, headers=h, timeout=30, **kw)
        try:
            return r.status_code, r.json()
        except ValueError:
            return r.status_code, r.text

    def login(self, u, p):
        s, j = self.req('POST', '/api/auth/login', json={'username': u, 'password': p})
        assert s == 200, j
        self.tok = j['token']
        return self

    get = lambda self, p, **kw: self.req('GET', p, **kw)
    post = lambda self, p, j=None: self.req('POST', p, json=j or {})
    put = lambda self, p, j=None: self.req('PUT', p, json=j or {})
    delete = lambda self, p: self.req('DELETE', p)


def main():
    tmp = tempfile.mkdtemp()
    port = socket.socket(); port.bind(('127.0.0.1', 0)); p = port.getsockname()[1]; port.close()
    env = dict(os.environ, DATABASE_URL=os.getenv('TEST_DATABASE_URL') or f'sqlite:///{tmp}/t.db', SECRET_KEY='x' * 48, SEED_LEGACY='1')
    srv = subprocess.Popen([sys.executable, '-m', 'uvicorn', 'app.main:app', '--port', str(p)], cwd=ROOT, env=env,
                           stdout=subprocess.DEVNULL, stderr=open(os.path.join(tmp, 'server.log'), 'w'))
    print('Log máy chủ:', os.path.join(tmp, 'server.log'))
    base = f'http://127.0.0.1:{p}'
    try:
        for _ in range(100):
            try:
                requests.get(base + '/api/health', timeout=1); break
            except requests.ConnectionError:
                time.sleep(0.2)
        run(base)
    finally:
        srv.terminate()
    print(f'\n{len(PASS)} đạt, {len(FAIL)} lỗi')
    sys.exit(1 if FAIL else 0)


def run(base):
    admin = Client(base).login('admin', 'Admin@123')
    print('1. Dữ liệu cũ')
    d = admin.get('/api/reports/dashboard?month=2026-05&date=2026-05-31')[1]
    check('Tổng số dư sau dữ liệu cũ = 863.233.490', d['kpi']['total'] == 863233490, d['kpi'])
    check('Kiểm soát hạch toán khớp', admin.get('/api/reports/control')[1]['ok'])

    opts = admin.get('/api/options')[1]
    co = {c['code']: c['id'] for c in opts['companies']}
    acc = {a['code']: a['id'] for a in opts['accounts']}

    print('2. Tài khoản người dùng')
    users = {
        'kt1': dict(role='staff', permissions=['cashbook.create', 'partner_payment.create', 'cash_fund.create'], company_ids=[co['VHS'], co['DP']]),
        'ktt': dict(role='chief_accountant', company_ids=[co['VHS'], co['DP'], co['MPN']]),
        'ktt2': dict(role='chief_accountant', company_ids=[co['TPE']]),
        'hr': dict(role='staff', permissions=['payroll.create'], company_ids=[co['VHS']]),
        'ktl': dict(role='staff', permissions=['payroll.approve'], company_ids=[co['VHS']]),
        'gdvhs': dict(role='director', company_ids=[co['VHS'], co['DP'], co['MPN'], co['VHS-LS']]),
    }
    for name, u in users.items():
        s, j = admin.post('/api/users', dict(u, username=name, full_name=name.upper(), password='123456'))
        check(f'Tạo tài khoản {name}', s == 201, j)
    s, j = admin.post('/api/users', dict(username='x', full_name='X', role='staff', password='123456'))
    check('Nhân viên bắt buộc có công ty', s == 422, j)
    kt1, ktt, ktt2, hr, ktl, gd = (Client(base).login(n, '123456') for n in users)

    print('3. Quy trình Thu Chi: Kế toán lập → KTT → GĐ công ty trụ cột → Admin tập đoàn')
    doc = dict(doc_date='2026-06-02', kind='Thu tiền khách hàng', account_id=acc['MB-VHS'], amount=50_000_000,
               counterparty='Khách A', description='Thu tiền hàng')
    s, j = kt1.post('/api/docs/cashbook', dict(doc, account_id=acc['MB-TPE']))
    check('Không lập được cho công ty ngoài phạm vi', s in (403, 422), j)
    s, j = kt1.post('/api/docs/cashbook', dict(doc, amount=0))
    check('Validate số tiền > 0', s == 422, j)
    s, j = hr.post('/api/docs/cashbook', doc)
    check('HR không có quyền lập Thu Chi', s == 403, j)
    s, j = kt1.post('/api/docs/cashbook', doc); did = j['id']
    check('Kế toán lập nháp', s == 201, j)
    check('Sửa nháp', kt1.put(f'/api/docs/cashbook/{did}', dict(doc, amount=60_000_000))[0] == 200)
    check('Gửi duyệt', kt1.post(f'/api/docs/cashbook/{did}/submit')[0] == 200)
    s, j = kt1.put(f'/api/docs/cashbook/{did}', doc)
    check('Không sửa được sau khi gửi', s in (400, 403, 409), j)
    s, j = kt1.post(f'/api/docs/cashbook/{did}/approve')
    check('Người lập không tự duyệt', s == 403, j)
    s, j = ktt2.post(f'/api/docs/cashbook/{did}/approve')
    check('KTT khác công ty không duyệt được', s == 403, j)
    s, j = ktl.post(f'/api/docs/cashbook/{did}/approve')
    check('Kế toán lương không duyệt Thu Chi', s == 403, j)
    check('KTT thấy trong hộp duyệt', any(x['id'] == did for x in ktt.get('/api/inbox')[1]))
    check('KTT duyệt cấp 1', ktt.post(f'/api/docs/cashbook/{did}/approve')[0] == 200)
    s, j = ktt.post(f'/api/docs/cashbook/{did}/reopen')
    check('KTT không mở lại được', s == 403, j)
    led = admin.get(f'/api/reports/ledger?from=2026-06-02&to=2026-06-02&account_id={acc["MB-VHS"]}')[1]
    check('Chưa ghi sổ sau duyệt cấp 1', led['inflow'] == 0, led)
    s, j = admin.post(f'/api/docs/cashbook/{did}/approve')
    check('Admin chưa duyệt được khi GĐ trụ cột chưa duyệt', s == 403, j)
    check('Trạng thái chờ GĐ trụ cột', admin.get(f'/api/docs/cashbook/{did}')[1]['status_label'] == 'Chờ GĐ công ty trụ cột')
    check('GĐ trụ cột thấy trong hộp duyệt', any(x['id'] == did and x['module'] == 'cashbook' for x in gd.get('/api/inbox')[1]))
    check('Admin chưa thấy trong hộp duyệt', not any(x['id'] == did and x['module'] == 'cashbook' for x in admin.get('/api/inbox')[1]))
    s, j = gd.post(f'/api/docs/cashbook/{did}/approve')
    check('GĐ trụ cột duyệt cấp 2', s == 200 and j['result'] == 'dir_approved', j)
    s, j = gd.post(f'/api/docs/cashbook/{did}/reopen')
    check('GĐ trụ cột không mở lại được', s == 403, j)
    led = admin.get(f'/api/reports/ledger?from=2026-06-02&to=2026-06-02&account_id={acc["MB-VHS"]}')[1]
    check('Chưa ghi sổ sau duyệt cấp 2', led['inflow'] == 0, led)
    check('Admin thấy trong hộp duyệt', any(x['id'] == did and x['module'] == 'cashbook' for x in admin.get('/api/inbox')[1]))
    check('Admin tập đoàn duyệt cuối', admin.post(f'/api/docs/cashbook/{did}/approve')[0] == 200)
    led = admin.get(f'/api/reports/ledger?from=2026-06-02&to=2026-06-02&account_id={acc["MB-VHS"]}')[1]
    check('Đã ghi sổ 60.000.000', led['inflow'] == 60_000_000, led)

    print('4. Mở lại chứng từ')
    s, j = kt1.post(f'/api/docs/cashbook/{did}/request_reopen', {})
    check('Đề nghị mở lại phải có lý do', s == 422 or s == 400, j)
    check('Đề nghị mở lại', kt1.post(f'/api/docs/cashbook/{did}/request_reopen', {'note': 'Sai số tiền'})[0] == 200)
    check('Giám đốc mở lại', admin.post(f'/api/docs/cashbook/{did}/reopen')[0] == 200)
    led = admin.get(f'/api/reports/ledger?from=2026-06-02&to=2026-06-02&account_id={acc["MB-VHS"]}')[1]
    check('Mở lại → hủy bút toán', led['inflow'] == 0, led)
    check('Người lập sửa lại', kt1.put(f'/api/docs/cashbook/{did}', dict(doc, _submit=True))[0] == 200)
    ktt.post(f'/api/docs/cashbook/{did}/approve'); gd.post(f'/api/docs/cashbook/{did}/approve'); admin.post(f'/api/docs/cashbook/{did}/approve')
    check('Duyệt lại → trạng thái posted', admin.get(f'/api/docs/cashbook/{did}')[1]['status'] == 'posted')

    print('5. Quỹ TM: tạm ứng NH → quỹ TM, chi tạm ứng NV')
    s, j = kt1.post('/api/docs/cash_fund', dict(doc_date='2026-06-03', kind='Tạm ứng quỹ TM', from_account_id=acc['MB-VHS'],
                                                to_account_id=acc['TM-VHS'], amount=10_000_000, _submit=True))
    check('Lập tạm ứng quỹ', s == 201, j); fid = j.get('id')
    s, j = kt1.post('/api/docs/cash_fund', dict(doc_date='2026-06-03', kind='Tạm ứng quỹ TM', from_account_id=acc['TM-VHS'],
                                                to_account_id=acc['MB-VHS'], amount=1))
    check('Sai loại TK bị chặn', s == 422, j)
    s, j = ktt.post('/api/docs/cash_fund/bulk', {'ids': [fid], 'action': 'approve'})
    check('Duyệt hàng loạt', j.get('done') == 1, j)
    check('GĐ trụ cột duyệt hàng loạt', gd.post('/api/docs/cash_fund/bulk', {'ids': [fid], 'action': 'approve'})[1].get('done') == 1)
    admin.post('/api/docs/cash_fund/bulk', {'ids': [fid], 'action': 'approve'})
    d = admin.get('/api/reports/summary-day?date=2026-06-03')[1]
    tm = next(a for a in d['accounts'] if a['code'] == 'TM-VHS')
    check('Quỹ TM VHS tăng 10.000.000', tm['closing'] == 10_000_000, tm)

    print('6. Lương: HR lập → Kế toán lương → GĐ trụ cột → Admin')
    pay = dict(period='2026-06', employee='Nguyễn Văn A', company_id=co['VHS'], gross=20_000_000, bonus=2_000_000,
               insurance=2_100_000, pit=500_000, doc_date='2026-06-05', account_id=acc['MB-VHS'])
    s, j = hr.post('/api/docs/payroll', dict(pay, _submit=True)); pid = j.get('id')
    check('HR lập bảng lương', s == 201, j)
    s, j = ktt.post(f'/api/docs/payroll/{pid}/approve')
    check('KTT không duyệt lương (không có quyền)', s == 403, j)
    check('Kế toán lương duyệt', ktl.post(f'/api/docs/payroll/{pid}/approve')[0] == 200)
    check('GĐ trụ cột duyệt lương', gd.post(f'/api/docs/payroll/{pid}/approve')[0] == 200)
    check('Admin duyệt lương', admin.post(f'/api/docs/payroll/{pid}/approve')[0] == 200)
    led = admin.get(f'/api/reports/ledger?from=2026-06-05&to=2026-06-05&account_id={acc["MB-VHS"]}')[1]
    check('Ghi sổ thực lĩnh 19.400.000', led['outflow'] == 19_400_000, led)

    print('7. Quyền xem')
    check('Nhân viên không xem tổng hợp', kt1.get('/api/reports/summary-month?month=2026-06')[0] == 403)
    check('KTT xem tổng hợp', ktt.get('/api/reports/summary-month?month=2026-06')[0] == 200)
    check('KTT không xem Dashboard', ktt.get('/api/reports/dashboard')[0] == 403)
    check('HR không xem Thu Chi', hr.get('/api/docs/cashbook')[0] == 403)
    s, j = kt1.get('/api/reports/ledger?from=2026-04-01&to=2026-06-30')
    check('Sổ chi tiết chỉ gồm công ty phụ trách', s == 200 and {r['company_code'] for r in j['rows']} <= {'VHS', 'DP'}, j.get('rows', [])[:1])
    check('Nhân viên không quản trị tài khoản', kt1.get('/api/users')[0] == 403)

    print('7b. GĐ công ty trụ cột: Dashboard theo nhóm')
    vhs_group = next(c['group_id'] for c in opts['companies'] if c['code'] == 'VHS')
    tpe_group = next(c['group_id'] for c in opts['companies'] if c['code'] == 'TPE')
    s, dg = gd.get('/api/reports/dashboard?month=2026-06&date=2026-06-30')
    check('GĐ trụ cột xem Dashboard', s == 200, dg)
    check('Dashboard GĐ trụ cột chỉ gồm công ty trong nhóm', {c['code'] for c in dg['companies']} <= {'VHS', 'DP', 'MPN', 'VHS-LS'}, [c['code'] for c in dg['companies']])
    s, da = admin.get(f'/api/reports/dashboard?month=2026-06&date=2026-06-30&group_id={vhs_group}')
    check('Admin lọc theo nhóm = số liệu GĐ trụ cột', da['kpi']['total'] == dg['kpi']['total'], (da['kpi']['total'], dg['kpi']['total']))
    s, dall = admin.get('/api/reports/dashboard?month=2026-06&date=2026-06-30')
    check('Admin xem toàn tập đoàn (lớn hơn 1 nhóm)', dall['kpi']['total'] > dg['kpi']['total'] and len(dall['group_options']) == 3, dall['group_options'])
    check('GĐ trụ cột không xem nhóm khác', gd.get(f'/api/reports/dashboard?group_id={tpe_group}')[0] == 403)
    check('GĐ trụ cột xem tổng hợp', gd.get('/api/reports/summary-month?month=2026-06')[0] == 200)
    check('GĐ trụ cột không quản trị tài khoản', gd.get('/api/users')[0] == 403)
    check('GĐ trụ cột không xem Kiểm soát', gd.get('/api/reports/control')[0] == 403)

    print('7c. Nhóm chưa có GĐ trụ cột (TPE): KTT duyệt xong chuyển thẳng Admin')
    s, j = admin.post('/api/docs/cashbook', dict(doc_date='2026-06-07', kind='Thu khác', account_id=acc['MB-TPE'], amount=1_000_000,
                                                description='Thu khác TPE', _submit=True)); tid = j.get('id')
    check('Admin lập chứng từ TPE', s == 201, j)
    check('KTT TPE duyệt cấp 1', ktt2.post(f'/api/docs/cashbook/{tid}/approve')[0] == 200)
    check('Trạng thái chờ Admin', admin.get(f'/api/docs/cashbook/{tid}')[1]['status_label'] == 'Chờ Admin duyệt cuối')
    check('Admin duyệt cuối ngay', admin.post(f'/api/docs/cashbook/{tid}/approve')[1].get('result') == 'posted')

    print('7d. Nhập bảng lương từ Excel')
    import io
    from openpyxl import load_workbook
    r = requests.get(base + '/api/docs/payroll/template', headers={'Authorization': f'Bearer {hr.tok}'})
    check('Tải biểu mẫu Excel', r.status_code == 200 and r.content[:2] == b'PK', r.status_code)
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb['Nhập liệu']
    heads = [c.value for c in ws[1]]
    check('Biểu mẫu có đủ cột', 'Kỳ lương *' in heads and 'Lương + phụ cấp *' in heads and 'TK chi trả *' in heads, heads)
    check('Danh mục chỉ gồm công ty của HR', [c.value for c in wb['Danh mục']['A'][1:] if c.value] == ['VHS'])
    col = {h.replace(' *', ''): i + 1 for i, h in enumerate(heads)}
    rows = [('07/2026', 'VHS', 'Trần Thị B', 'Kế toán', 15_000_000, 1_000_000, 1_575_000, 0, 0, 0, '05/08/2026', 'MB-VHS'),
            ('07/2026', 'VHS', 'Lê Văn C', 'Kho', 12_000_000, None, 1_260_000, None, 500_000, None, '05/08/2026', 'MB-VHS'),
            ('07/2026', 'VHS', 'Phạm D', 'Kho', 9_000_000, None, None, None, None, None, '05/08/2026', 'SAI-TK')]
    keys = ['Kỳ lương', 'Công ty hạch toán lương', 'Nhân viên', 'Bộ phận', 'Lương + phụ cấp', 'Thưởng', 'BHXH, BHYT (NLĐ)', 'Thuế TNCN',
            'Trừ tạm ứng', 'Khấu trừ khác', 'Ngày chi trả', 'TK chi trả']
    for rno, vals in enumerate(rows, 2):
        for k, v in zip(keys, vals):
            ws.cell(row=rno, column=col[k], value=v)
    ws.cell(row=5, column=col['Kỳ lương'], value='07/2026'); ws.cell(row=5, column=col['Công ty hạch toán lương'], value='VHS')
    ws.cell(row=5, column=col['Nhân viên'], value='Trần Thị B'); ws.cell(row=5, column=col['Lương + phụ cấp'], value=1)
    ws.cell(row=5, column=col['Ngày chi trả'], value='05/08/2026'); ws.cell(row=5, column=col['TK chi trả'], value='MB-VHS')
    buf = io.BytesIO(); wb.save(buf)
    import base64
    b64 = base64.b64encode(buf.getvalue()).decode()
    s, j = hr.post('/api/docs/payroll/import', {'file': b64, 'filename': 'luong.xlsx', 'dry_run': True})
    check('Xem trước: 4 dòng, 2 lỗi (sai TK, trùng nhân viên)', s == 200 and j['total'] == 4 and j['errors'] == 2, j if s != 200 else [r['errors'] for r in j['rows']])
    check('Thực lĩnh dòng 1 = 14.425.000', j['rows'][0]['amount'] == 14_425_000, j['rows'][0])
    s, j = hr.post('/api/docs/payroll/import', {'file': b64, 'filename': 'luong.xlsx', 'dry_run': False, 'submit': True})
    check('Còn lỗi → không nhập dòng nào', s == 422, j)
    check('Chưa có chứng từ nào được tạo', hr.get('/api/docs/payroll?q=Trần Thị B')[1]['total'] == 0)
    ws.delete_rows(4, 2); buf = io.BytesIO(); wb.save(buf); b64 = base64.b64encode(buf.getvalue()).decode()
    s, j = hr.post('/api/docs/payroll/import', {'file': b64, 'filename': 'luong.xlsx', 'dry_run': False, 'submit': True})
    check('Nhập 2 chứng từ & gửi duyệt', s == 201 and j['created'] == 2, j)
    lst = ktl.get('/api/docs/payroll?status=submitted')[1]
    check('Kế toán lương thấy 2 chứng từ chờ duyệt', lst['total'] == 2, lst['total'])
    s, j = hr.post('/api/docs/payroll/import', {'file': b64, 'filename': 'luong.xlsx', 'dry_run': True})
    check('Nhập lại cùng file → báo trùng', j['errors'] == 2, j['errors'])
    check('Kế toán (không có quyền lương) không tải biểu mẫu', requests.get(base + '/api/docs/payroll/template', headers={'Authorization': f'Bearer {kt1.tok}'}).status_code == 403)

    print('8. Dữ liệu khai báo: khóa / mở khóa')
    s, j = kt1.post('/api/master/partners', dict(name='NCC Test', type='supplier'))
    check('Nhân viên thêm đối tác', s == 201, j); pid = j.get('id')
    check('Sửa khi chưa khóa', kt1.put(f'/api/master/partners/{pid}', dict(name='NCC Test 2', type='supplier'))[0] == 200)
    check('Khóa (submit)', kt1.post('/api/master/partners/lock', {'ids': [pid]})[0] == 200)
    check('Không sửa sau khóa', kt1.put(f'/api/master/partners/{pid}', dict(name='X', type='supplier'))[0] in (403, 409))
    check('Không xóa sau khóa', kt1.delete(f'/api/master/partners/{pid}')[0] in (403, 409))
    check('KTT không mở khóa', ktt.post('/api/master/partners/unlock', {'ids': [pid]})[0] == 403)
    check('GĐ mở khóa', admin.post('/api/master/partners/unlock', {'ids': [pid]})[0] == 200)
    check('Xóa sau khi mở khóa', kt1.delete(f'/api/master/partners/{pid}')[0] == 200)

    print('9. Hủy & kiểm soát')
    s, j = kt1.post('/api/docs/partner_payment', dict(doc_date='2026-06-06', account_id=acc['MB-VHS'], partner='NCC B',
                                                      amount=5_000_000, description='Thanh toán hàng', _submit=True))
    check('Lập thanh toán đối tác', s == 201, j)
    check('GĐ hủy chứng từ', admin.post(f'/api/docs/partner_payment/{j["id"]}/cancel', {'note': 'Trùng'})[0] == 200)
    check('Kiểm soát hạch toán vẫn khớp', admin.get('/api/reports/control')[1]['ok'])
    check('Xóa user có chứng từ → khóa', admin.delete(f'/api/users/{admin.get("/api/users")[1][1]["id"]}')[1]['result'] == 'deactivated')


if __name__ == '__main__':
    main()
