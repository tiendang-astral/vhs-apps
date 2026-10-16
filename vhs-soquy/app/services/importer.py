"""Nhập chứng từ hàng loạt từ biểu mẫu Excel (dùng chung cho mọi nghiệp vụ có `importable = True`, hiện tại: Thanh toán lương).

- template(): sinh file .xlsx gồm sheet nhập liệu (cột theo trường của nghiệp vụ, có danh sách chọn công ty / TK),
  sheet Danh mục (mã công ty, mã TK, nhân viên trong phạm vi người dùng) và sheet Hướng dẫn.
- preview(): đọc file, chuyển mã → id, kiểm tra từng dòng y như nhập tay + phát hiện trùng (trong file và đã có trên hệ thống).
- commit(): chỉ khi KHÔNG còn dòng lỗi mới tạo chứng từ (tất cả trong 1 transaction – hoặc nhập hết, hoặc không nhập dòng nào).
"""
import base64
import datetime as dt
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..http import ApiError
from ..modules.base import Ctx, to_float, to_int
from ..permissions import can_create, require
from . import workflow as wf

MAX_ROWS = 3000
MAX_BYTES = 5 * 1024 * 1024
HINT = {'month': 'MM/YYYY, VD 09/2026', 'date': 'dd/mm/yyyy', 'money': 'Số, không âm', 'company': 'Mã công ty (chọn từ danh sách)',
        'account': 'Mã TK (chọn từ danh sách)', 'category': 'Tên hạng mục', 'partner': 'Họ tên đúng như Dữ liệu khai báo'}
BLUE, GREY = '1F4FD8', 'F2F5FB'


def _fields(m):
    names = m.import_fields or [f.name for f in m.fields if not f.computed]
    by = {f.name: f for f in m.fields}
    return [by[n] for n in names]


def _header(f):
    return f.label + (' *' if f.required else '')


def _norm(s):
    s = re.sub(r'\(.*?\)', '', str(s or '')).replace('*', '')
    return re.sub(r'\s+', ' ', s).strip().lower()


def _scope_lists(db, user, m):
    mine = set(user['company_ids'])
    companies = [c for c in db.all("SELECT id, code, name FROM companies WHERE status='active' ORDER BY id") if c['id'] in mine]
    any_acc = any(f.type == 'account' and f.source == 'any' for f in m.fields)
    accounts = [a for a in db.all("SELECT a.id, a.code, a.kind, c.code AS company_code, a.company_id FROM accounts a "
                                  "JOIN companies c ON c.id=a.company_id WHERE a.status='active' ORDER BY c.id, a.kind, a.code")
                if any_acc or a['company_id'] in mine]
    people = [p for p in db.all("SELECT p.name, c.code AS company_code, p.company_id FROM partners p "
                                "LEFT JOIN companies c ON c.id=p.company_id WHERE p.type='employee' ORDER BY p.name")
              if p['company_id'] is None or p['company_id'] in mine]
    return companies, accounts, people


# ================================================================ biểu mẫu
def template(db, user, m):
    require(m.importable and can_create(user, m), 'Bạn không có quyền nhập nghiệp vụ này')
    fields = _fields(m)
    companies, accounts, people = _scope_lists(db, user, m)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Nhập liệu'
    thin = Side(style='thin', color='D0D7E2')
    head_fill, head_font = PatternFill('solid', fgColor=BLUE), Font(bold=True, color='FFFFFF')
    for i, f in enumerate(fields, 1):
        c = ws.cell(row=1, column=i, value=_header(f))
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = Border(bottom=thin, right=thin)
        hint = HINT.get(f.type, '')
        if f.help:
            hint = (hint + '. ' if hint else '') + f.help
        if hint:
            c.comment = Comment(hint, 'Hệ thống')
        col = get_column_letter(i)
        ws.column_dimensions[col].width = 16 if f.type in ('money', 'date', 'month', 'company') else 24 if f.type != 'textarea' else 30
        for r in range(2, 1002):
            cell = ws.cell(row=r, column=i)
            if f.type == 'money':
                cell.number_format = '#,##0'
            elif f.type == 'date':
                cell.number_format = 'dd/mm/yyyy'
            elif f.type == 'month':
                cell.number_format = '@'
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = 'A2'

    # Danh mục để chọn
    ref = wb.create_sheet('Danh mục')
    blocks = [('Mã công ty', 'Tên công ty', [(c['code'], c['name']) for c in companies]),
              ('Mã TK chi trả', 'Công ty · Loại', [(a['code'], f"{a['company_code']} · {'Tiền mặt' if a['kind'] == 'cash' else 'Ngân hàng'}") for a in accounts]),
              ('Nhân viên', 'Công ty', [(p['name'], p['company_code'] or '') for p in people])]
    for b, (h1, h2, rows) in enumerate(blocks):
        c0 = b * 3 + 1
        for j, h in enumerate((h1, h2)):
            cell = ref.cell(row=1, column=c0 + j, value=h)
            cell.fill, cell.font = head_fill, head_font
            ref.column_dimensions[get_column_letter(c0 + j)].width = 22 if j == 0 else 40
        for r, (v1, v2) in enumerate(rows, 2):
            ref.cell(row=r, column=c0, value=v1)
            ref.cell(row=r, column=c0 + 1, value=v2)
    lists = {'company': ('A', len(companies)), 'account': ('D', len(accounts)), 'employee': ('G', len(people))}
    for i, f in enumerate(fields, 1):
        key = 'employee' if f.type == 'partner' and f.source == 'employee' else f.type
        if key in lists and lists[key][1]:
            colref, n = lists[key]
            dv = DataValidation(type='list', formula1=f"='Danh mục'!${colref}$2:${colref}${n + 1}", allow_blank=True)
            if key == 'employee':                    # cho phép gõ tên chưa có trong danh mục
                dv.showErrorMessage = False
            else:
                dv.error, dv.errorTitle = 'Chọn mã có trong sheet Danh mục', 'Giá trị không hợp lệ'
            ws.add_data_validation(dv)
            dv.add(f'{get_column_letter(i)}2:{get_column_letter(i)}1001')
        if f.type == 'select' and f.options:
            dv = DataValidation(type='list', formula1='"' + ','.join(f.options) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f'{get_column_letter(i)}2:{get_column_letter(i)}1001')

    # Hướng dẫn + dòng ví dụ
    g = wb.create_sheet('Hướng dẫn')
    g.column_dimensions['A'].width = 110
    lines = [f'BIỂU MẪU NHẬP {m.title.upper()}', '',
             '1. Điền dữ liệu vào sheet "Nhập liệu", mỗi dòng là 1 chứng từ (với lương: 1 nhân viên / 1 kỳ lương). Không đổi tên cột.',
             '2. Cột có dấu * là bắt buộc. Mã công ty và Mã TK chọn từ danh sách (lấy từ sheet "Danh mục").',
             '3. Số tiền nhập dạng số (VD 15000000). Ngày dạng dd/mm/yyyy. Kỳ lương dạng MM/YYYY (VD 09/2026).',
             '4. Thực lĩnh hệ thống tự tính = Lương + phụ cấp + Thưởng − BHXH − Thuế TNCN − Trừ tạm ứng − Khấu trừ khác.'
             if m.key == 'payroll' else '4. Hệ thống tự kiểm tra dữ liệu giống như khi nhập tay.',
             '5. Trên phần mềm: bấm "Nhập từ Excel" → chọn file → xem kết quả kiểm tra từng dòng → bấm Nhập.',
             '   Nếu còn dòng lỗi, hệ thống KHÔNG nhập dòng nào; sửa file theo thông báo rồi tải lên lại.',
             '6. Dòng trùng (cùng công ty + kỳ lương + nhân viên đã có trên hệ thống hoặc lặp trong file) sẽ bị báo lỗi.',
             '', 'Các cột:']
    for f in fields:
        lines.append(f'   • {_header(f)}: {HINT.get(f.type, "Chữ")}' + (f' – {f.help}' if f.help else ''))
    for r, t in enumerate(lines, 1):
        cell = g.cell(row=r, column=1, value=t)
        if r == 1:
            cell.font = Font(bold=True, size=14, color=BLUE)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ================================================================ đọc & kiểm tra
def _date(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    m = re.match(r'^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})$', s)
    if m:
        return f'{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}'
    if re.match(r'^\d{4}-\d{2}-\d{2}', s):
        return s[:10]
    return '!' + s


def _month(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return f'{v.year:04d}-{v.month:02d}'
    s = str(v).strip()
    m = re.match(r'^(\d{1,2})[/.-](\d{4})$', s)
    if m:
        return f'{m[2]}-{int(m[1]):02d}'
    m = re.match(r'^(\d{4})-(\d{1,2})', s)
    if m:
        return f'{m[1]}-{int(m[2]):02d}'
    return '!' + s


def _text(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _read(content):
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        raise ApiError(400, 'Không đọc được file. Vui lòng dùng file .xlsx tải từ nút "Tải biểu mẫu Excel".')
    ws = wb['Nhập liệu'] if 'Nhập liệu' in wb.sheetnames else wb.worksheets[0]
    return list(ws.iter_rows(values_only=True))


def preview(db, user, m, content):
    require(m.importable and can_create(user, m), 'Bạn không có quyền nhập nghiệp vụ này')
    rows = _read(content)
    if not rows:
        raise ApiError(400, 'File không có dữ liệu')
    fields = _fields(m)
    head = [_norm(h) for h in rows[0]]
    colmap, missing = {}, []
    for f in fields:
        keys = {_norm(f.label), f.name.lower()}
        idx = next((i for i, h in enumerate(head) if h in keys), None)
        if idx is None:
            if f.required:
                missing.append(f.label)
        else:
            colmap[f.name] = idx
    if missing:
        raise ApiError(400, 'File thiếu cột: ' + ', '.join(missing) + '. Vui lòng dùng đúng biểu mẫu.')

    ctx = Ctx(db)
    co_by_code = {r['code'].upper(): r['id'] for r in db.all('SELECT id, code FROM companies')}
    co_code = {v: k for k, v in co_by_code.items()}
    acc_by_code = {a['code'].upper(): i for i, a in ctx.accounts.items()}
    existing = set()
    for d in db.all(f"SELECT * FROM {m.table} WHERE status<>'cancelled'"):
        k = m.duplicate_key(d)
        if k:
            existing.add(k)
    seen, out = {}, []
    labels = {f.name: f.label for f in m.fields}

    for rno, row in enumerate(rows[1:], 2):
        if row is None or all(v in (None, '') for v in row):
            continue
        if len(out) >= MAX_ROWS:
            raise ApiError(400, f'Mỗi lần nhập tối đa {MAX_ROWS} dòng')
        data, err = {}, {}
        for f in fields:
            if f.name not in colmap:
                continue
            v = row[colmap[f.name]] if colmap[f.name] < len(row) else None
            if v in (None, ''):
                data[f.name] = None
                continue
            if f.type == 'money':
                data[f.name] = to_int(v)
            elif f.type == 'number':
                data[f.name] = to_float(v)
            elif f.type == 'date':
                data[f.name] = _date(v)
            elif f.type == 'month':
                data[f.name] = _month(v)
            elif f.type == 'company':
                code = _text(v).split(' ')[0].upper()
                data[f.name] = co_by_code.get(code)
                if not data[f.name]:
                    err[f.name] = f'Mã công ty "{_text(v)}" không tồn tại'
            elif f.type == 'account':
                data[f.name] = acc_by_code.get(_text(v).upper())
                if not data[f.name]:
                    err[f.name] = f'Mã TK "{_text(v)}" không tồn tại'
            elif f.type == 'category':
                data[f.name] = ctx.cat_id(_text(v))
                if not data[f.name]:
                    err[f.name] = f'Hạng mục "{_text(v)}" không tồn tại'
            else:
                data[f.name] = _text(v)
            if isinstance(data[f.name], str) and data[f.name].startswith('!') and f.type in ('date', 'month'):
                err[f.name] = f'"{data[f.name][1:]}" không đúng định dạng {HINT[f.type]}'
                data[f.name] = None
        doc = m.normalize(data)
        for k, v in m.validate(doc, ctx).items():
            err.setdefault(k, v)
        company = m.company_of(doc, ctx)
        if company and company not in user['company_ids']:
            err.setdefault('company_id', f'Bạn không phụ trách công ty {co_code.get(company, company)}')
        key = m.duplicate_key(dict(doc, company_id=company))
        if key and not err:
            if key in existing:
                err['_'] = 'Trùng: đã có chứng từ cùng công ty / kỳ / nhân viên trên hệ thống'
            elif key in seen:
                err['_'] = f'Trùng với dòng {seen[key]} trong file'
            else:
                seen[key] = rno
        view = {}
        for f in fields:
            v = doc.get(f.name)
            if f.type == 'company':
                v = co_code.get(v, '')
            elif f.type == 'account':
                v = ctx.accounts.get(v, {}).get('code', '')
            elif f.type == 'month' and v:
                v = f'{v[5:7]}/{v[:4]}'
            view[f.name] = v
        out.append({'row': rno, 'data': data, 'view': view, 'amount': m.amount(doc) if not err else m.amount(doc),
                    'errors': {(labels.get(k, 'Dòng')): v for k, v in err.items()}})
    if not out:
        raise ApiError(400, 'File chưa có dòng dữ liệu nào (dòng 1 là tiêu đề cột)')
    bad = sum(1 for r in out if r['errors'])
    return {'rows': out, 'total': len(out), 'errors': bad, 'valid': len(out) - bad,
            'amount': sum(r['amount'] or 0 for r in out if not r['errors']),
            'columns': [{'name': f.name, 'label': f.label, 'type': f.type} for f in fields]}


def decode(b64):
    try:
        content = base64.b64decode((b64 or '').split(',')[-1], validate=False)
    except Exception:
        raise ApiError(400, 'File không hợp lệ')
    if not content:
        raise ApiError(400, 'Chưa chọn file')
    if len(content) > MAX_BYTES:
        raise ApiError(400, 'File quá lớn (tối đa 5 MB)')
    return content


def commit(db, user, m, content, filename='', submit=False):
    res = preview(db, user, m, content)
    if res['errors']:
        raise ApiError(422, f"Còn {res['errors']} dòng lỗi – chưa nhập dòng nào. Vui lòng sửa file rồi tải lên lại.")
    ids = []
    for r in res['rows']:
        doc_id = wf.create(db, user, m, r['data'])
        wf.log(db, m, doc_id, 'Nhập từ file Excel', user, f"{filename} – dòng {r['row']}" if filename else f"dòng {r['row']}")
        if submit:
            wf.submit(db, user, m, doc_id)
        ids.append(doc_id)
    return {'created': len(ids), 'ids': ids, 'submitted': bool(submit), 'amount': res['amount']}
